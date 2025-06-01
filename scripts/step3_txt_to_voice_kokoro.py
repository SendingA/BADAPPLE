from kokoro import KModel, KPipeline
import numpy as np
import soundfile as sf
import torch
import tqdm
import os
import os
import openpyxl
import asyncio
import tqdm
import argparse
import json
import re

class SpeechProvider:
    def __init__(self, gender, language):
        REPO_ID = 'hexgrad/Kokoro-82M-v1.1-zh'
        kororo_path = os.path.join(os.path.dirname(__file__), '..', 'voice', 'Kokoro-82M-v1.1-zh')
        config = kororo_path + '/config.json'
        model_pth = kororo_path + '/kokoro-v1_1-zh.pth'

        self.SAMPLE_RATE = 24000
        # How much silence to insert between paragraphs: 5000 is about 0.2 seconds
        self.N_ZEROS = 10000
        # VOICES = glob.glob(f"./voice/Kokoro-82M-v1.1-zh/voices/{GENDER}*.pt")
        self.VOICE = kororo_path + '/voices/zf_003.pt' if gender == 'zf' else kororo_path + '/voices/zm_031.pt'
        self.device = 'cuda' if torch.cuda.is_available() else 'cpu'
        self.model = KModel(repo_id=REPO_ID, config=config, model=model_pth).to(self.device).eval()
        self.en_pipeline = KPipeline(lang_code='a', repo_id=REPO_ID, model=False)
        self.zh_pipeline = KPipeline(lang_code='z', repo_id=REPO_ID, model=self.model, en_callable=self.en_callable)
        self.zh_pipeline.load_voice


    def en_callable(self,text):
        if text == 'Kokoro':
            return 'kˈOkəɹO'
        elif text == 'Sol':
            return 'sˈOl'
        return next(self.en_pipeline(text)).phonemes

    def speed_callable(self,len_ps):
        speed = 1
        return speed 

    def get_tts_audio(self, message):
        wavs_tot = []
        durations = []  # 添加时长记录
        
        for paragraph in tqdm.tqdm(message, desc="正在生成配音音频", unit="paragraphs"):
            wavs_para = []
            sentence_durations = []  # 记录每个句子的时长
            
            for i, sentence in enumerate(paragraph):
                # print(f"正在处理句子: {sentence}")
                generator = self.zh_pipeline(sentence, voice=self.VOICE, speed=self.speed_callable)
                result = next(generator)
                wav = result.audio
                
                # 计算当前句子的时长（秒）
                sentence_duration = len(wav) / self.SAMPLE_RATE
                sentence_durations.append(sentence_duration)
                
                if i == 0 and wavs_tot and self.N_ZEROS > 0:
                    wav = np.concatenate([np.zeros(self.N_ZEROS), wav])
                    # 添加静音时长
                    sentence_durations[-1] += self.N_ZEROS / self.SAMPLE_RATE
                    
                if i == 0:
                    wavs_para = wav
                else:
                    wavs_para = np.concatenate([wavs_para, wav])
                    
            wavs_tot.append(wavs_para)
            durations.append(sentence_durations)

        return 'wav', wavs_tot, durations

def convert_text_to_audio(tasks, language, output_path, gender):
    if not tasks:
        return False

    provider = SpeechProvider(gender, language)
    wav_format, wavs, durations = provider.get_tts_audio(tasks)
    
    if wav_format != 'wav':
        raise ValueError("Unsupported audio format")
        
    # 保存音频文件和时长信息
    timing_info = {}
    audio_files = []
    if wavs is not None:
        for index, (wav_para, duration_list) in enumerate(zip(wavs, durations)):
            wav_file_path = os.path.join(output_path, f"output_{index}.wav")
            sf.write(wav_file_path, wav_para, provider.SAMPLE_RATE)
            audio_files.append(wav_file_path)
            
            # 记录时长信息
            total_duration = len(wav_para) / provider.SAMPLE_RATE
            timing_info[f"output_{index}"] = {
                "total_duration": total_duration,
                "sentence_durations": duration_list,
                "sample_rate": provider.SAMPLE_RATE
            }
    
    # 保存时长信息到JSON文件
    timing_file = os.path.join(output_path, "audio_timing.json")
    with open(timing_file, 'w', encoding='utf-8') as f:
        json.dump(timing_info, f, indent=2, ensure_ascii=False)
    
    print(f"✅ 音频时长信息已保存到: {timing_file}")
    return True, audio_files

def process_text_files(input_file, output_dir, language, gender):
    print("BADAPPLE")
    with open(input_file, 'r', encoding='utf-8') as f:
        data = json.load(f)
    scenarios = [item["内容"] for item in data.values()]
    tasks = []
    sentence_mapping = []  # 记录句子与原文的映射关系
    
    for scenario_idx, scenario in enumerate(scenarios):
        print("="*50)
        # 先处理双引号包裹的句子
        quote_pattern = r'“([^”]*[。！？])”'
        quotes = re.findall(quote_pattern, scenario)
        
        # 将双引号内容替换为占位符
        temp_scenario = scenario
        placeholders = {}
        for i, quote in enumerate(quotes):
            placeholder = f"__QUOTE_{i}__"
            placeholders[placeholder] = quote
            temp_scenario = temp_scenario.replace(f'“{quote}”', placeholder)

        sentences = re.split(r'([。！？.!?])', temp_scenario)
        processed_sentences = []
        
        for i in range(0, len(sentences), 2):
            if i < len(sentences):
                s = sentences[i].strip()
                if s:
                    # 添加标点符号（如果存在）
                    if i + 1 < len(sentences) and sentences[i + 1]:
                        s += sentences[i + 1]
                    for placeholder, quote in placeholders.items():
                        s = s.replace(placeholder, f'“{quote}”')
                    processed_sentences.append(s)

        if processed_sentences:
            tasks.append(tuple(processed_sentences))
            sentence_mapping.append({
                "scenario_index": scenario_idx,
                "original_text": scenario,
                "processed_sentences": processed_sentences
            })
    
    # 保存句子映射信息
    mapping_file = os.path.join(output_dir, "sentence_mapping.json")
    with open(mapping_file, 'w', encoding='utf-8') as f:
        json.dump(sentence_mapping, f, indent=2, ensure_ascii=False)
    
    print(f"✅ 句子映射信息已保存到: {mapping_file}")
    
    return convert_text_to_audio(tasks, language, output_dir, gender)


def main(input_file, output_dir, language, gender):
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)
    return process_text_files(input_file, output_dir, language, gender)


if __name__ == "__main__":
    main()
